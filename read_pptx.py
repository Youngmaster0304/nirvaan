from pptx import Presentation
from pptx.util import Inches, Pt
import json

prs = Presentation(r'C:\Users\Abhinav\Downloads\SIH2026-IDEA-Presentation-Format (1).pptx')

slide_data = []
for i, slide in enumerate(prs.slides):
    slide_info = {'slide_number': i+1, 'shapes': []}
    for shape in slide.shapes:
        shape_info = {
            'name': shape.name,
            'shape_type': str(shape.shape_type),
            'left': shape.left,
            'top': shape.top,
            'width': shape.width,
            'height': shape.height
        }
        if hasattr(shape, 'text') and shape.text:
            shape_info['text'] = shape.text[:300]
        if shape.has_text_frame:
            paragraphs = []
            for para in shape.text_frame.paragraphs:
                para_text = para.text
                if para_text:
                    paragraphs.append(para_text[:200])
            shape_info['paragraphs'] = paragraphs
        slide_info['shapes'].append(shape_info)
    slide_data.append(slide_info)

for slide in slide_data:
    print(f"=== Slide {slide['slide_number']} ===")
    for shape in slide['shapes']:
        if 'text' in shape and shape['text'].strip():
            print(f"  [{shape['name']}] {shape['text'][:200]}")
        if 'paragraphs' in shape:
            for p in shape['paragraphs']:
                if p.strip():
                    print(f"    - {p[:150]}")
    print()
