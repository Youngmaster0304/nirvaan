from pptx import Presentation
from pptx.util import Inches, Pt, Emu
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.enum.shapes import MSO_SHAPE
import os

prs = Presentation()
prs.slide_width = Inches(13.333)
prs.slide_height = Inches(7.5)

# Colors
PRIMARY = RGBColor(0x1a, 0x56, 0xdb)  # Blue
DARK = RGBColor(0x0f, 0x17, 0x2a)
WHITE = RGBColor(0xFF, 0xFF, 0xFF)
GRAY = RGBColor(0x94, 0xa3, 0xb8)
GREEN = RGBColor(0x10, 0xb9, 0x81)
ORANGE = RGBColor(0xf5, 0x9e, 0x0b)
RED = RGBColor(0xef, 0x44, 0x44)

def add_background(slide, color):
    background = slide.background
    fill = background.fill
    fill.solid()
    fill.fore_color.rgb = color

def add_textbox(slide, left, top, width, height, text, font_size=18, bold=False, color=WHITE, alignment=PP_ALIGN.LEFT):
    txBox = slide.shapes.add_textbox(left, top, width, height)
    tf = txBox.text_frame
    tf.word_wrap = True
    p = tf.paragraphs[0]
    p.text = text
    p.font.size = Pt(font_size)
    p.font.bold = bold
    p.font.color.rgb = color
    p.alignment = alignment
    return txBox

def add_bullet_text(slide, left, top, width, height, items, font_size=16, color=WHITE):
    txBox = slide.shapes.add_textbox(left, top, width, height)
    tf = txBox.text_frame
    tf.word_wrap = True
    for i, item in enumerate(items):
        if i == 0:
            p = tf.paragraphs[0]
        else:
            p = tf.add_paragraph()
        p.text = item
        p.font.size = Pt(font_size)
        p.font.color.rgb = color
        p.space_after = Pt(8)
        p.level = 0
    return txBox

def add_shape_box(slide, left, top, width, height, fill_color, border_color=None):
    shape = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, left, top, width, height)
    shape.fill.solid()
    shape.fill.fore_color.rgb = fill_color
    if border_color:
        shape.line.color.rgb = border_color
        shape.line.width = Pt(1)
    else:
        shape.line.fill.background()
    return shape

# ============================================
# SLIDE 1: TITLE PAGE
# ============================================
slide1 = prs.slides.add_slide(prs.slide_layouts[6])  # Blank
add_background(slide1, DARK)

# Title
add_textbox(slide1, Inches(1), Inches(0.8), Inches(11), Inches(0.8),
            "SMART INDIA HACKATHON 2026", font_size=20, bold=True, color=GRAY, alignment=PP_ALIGN.CENTER)

# Main Title
add_textbox(slide1, Inches(1), Inches(1.8), Inches(11), Inches(1.2),
            "Niravaan", font_size=54, bold=True, color=PRIMARY, alignment=PP_ALIGN.CENTER)

add_textbox(slide1, Inches(1), Inches(3.0), Inches(11), Inches(0.8),
            "AI-Powered Automatic Block Planning to Maximize\nAsset Availability for Train Operations on Indian Railways",
            font_size=22, color=WHITE, alignment=PP_ALIGN.CENTER)

# Details box
box = add_shape_box(slide1, Inches(3), Inches(4.2), Inches(7), Inches(2.5), RGBColor(0x1e, 0x29, 0x3b))
details = [
    "Problem Statement ID: SIH26027",
    "Problem Statement Title: AI-Powered Automatic Block Planning",
    "Theme: Transportation & Logistics",
    "PS Category: Software",
    "Team ID: [Your Team ID]",
    "Team Name: [Your Team Name]"
]
add_bullet_text(slide1, Inches(3.5), Inches(4.4), Inches(6), Inches(2.2), details, font_size=16, color=WHITE)

# ============================================
# SLIDE 2: IDEA TITLE - PROPOSED SOLUTION
# ============================================
slide2 = prs.slides.add_slide(prs.slide_layouts[6])
add_background(slide2, DARK)

add_textbox(slide2, Inches(0.5), Inches(0.3), Inches(12), Inches(0.6),
            "IDEA TITLE", font_size=28, bold=True, color=PRIMARY)

add_textbox(slide2, Inches(0.5), Inches(0.9), Inches(12), Inches(0.5),
            "Proposed Solution: Niravaan - Intelligent Block Planning System",
            font_size=20, color=WHITE)

# Left column - Problem
box1 = add_shape_box(slide2, Inches(0.5), Inches(1.6), Inches(6), Inches(5.5), RGBColor(0x1e, 0x29, 0x3b))
add_textbox(slide2, Inches(0.8), Inches(1.7), Inches(5.5), Inches(0.5),
            "Current Problems", font_size=18, bold=True, color=RED)
problems = [
    "Departments plan maintenance independently",
    "Data silos across TMS, SMMS, TDMS, COA",
    "Manual block requests via phone calls",
    "No audit trail for block refusals",
    "Maintenance corridors violated frequently",
    "Safety incidents due to delayed maintenance"
]
add_bullet_text(slide2, Inches(0.8), Inches(2.3), Inches(5.5), Inches(4.5), problems, font_size=14, color=GRAY)

# Right column - Solution
box2 = add_shape_box(slide2, Inches(6.8), Inches(1.6), Inches(6), Inches(5.5), RGBColor(0x1e, 0x29, 0x3b))
add_textbox(slide2, Inches(7.1), Inches(1.7), Inches(5.5), Inches(0.5),
            "Our Solution", font_size=18, bold=True, color=GREEN)
solutions = [
    "AI/ML engine for optimized block scheduling",
    "Real-time integration of all railway systems",
    "Automated multi-department coordination",
    "Weekly + monthly block plan generation",
    "Criticality-based defect prioritization",
    "Dashboard with AI recommendations"
]
add_bullet_text(slide2, Inches(7.1), Inches(2.3), Inches(5.5), Inches(4.5), solutions, font_size=14, color=GRAY)

# Innovation box
add_textbox(slide2, Inches(0.5), Inches(6.3), Inches(12), Inches(0.5),
            "Innovation: First system to use Genetic Algorithm + Reinforcement Learning for multi-department railway block optimization",
            font_size=14, bold=True, color=PRIMARY)

# ============================================
# SLIDE 3: TECHNICAL APPROACH
# ============================================
slide3 = prs.slides.add_slide(prs.slide_layouts[6])
add_background(slide3, DARK)

add_textbox(slide3, Inches(0.5), Inches(0.3), Inches(12), Inches(0.6),
            "TECHNICAL APPROACH", font_size=28, bold=True, color=PRIMARY)

# Tech Stack
add_textbox(slide3, Inches(0.5), Inches(1.0), Inches(12), Inches(0.4),
            "Technologies Used", font_size=18, bold=True, color=WHITE)

tech_items = [
    ("Frontend", "React.js, Chart.js, Tailwind CSS", PRIMARY),
    ("Backend", "Python, FastAPI, REST APIs", GREEN),
    ("AI/ML", "TensorFlow, Genetic Algorithm, RL, GNN", ORANGE),
    ("Database", "PostgreSQL, Redis (caching)", RED),
    ("Deployment", "Docker, Kubernetes, Cloud", GRAY),
]

for i, (title, desc, color) in enumerate(tech_items):
    x = Inches(0.5 + (i % 3) * 4.2)
    y = Inches(1.5 + (i // 3) * 1.2)
    box = add_shape_box(slide3, x, y, Inches(3.8), Inches(1.0), RGBColor(0x1e, 0x29, 0x3b), color)
    add_textbox(slide3, x + Inches(0.2), y + Inches(0.1), Inches(3.4), Inches(0.3),
                title, font_size=14, bold=True, color=color)
    add_textbox(slide3, x + Inches(0.2), y + Inches(0.4), Inches(3.4), Inches(0.5),
                desc, font_size=12, color=GRAY)

# Architecture Flow
add_textbox(slide3, Inches(0.5), Inches(3.8), Inches(12), Inches(0.4),
            "System Architecture Flow", font_size=18, bold=True, color=WHITE)

flow_steps = [
    ("Data Sources", "TMS + SMMS + TDMS + COA"),
    ("Integration", "API Layer + ETL"),
    ("AI Engine", "GA + RL + GNN"),
    ("Output", "Block Plans"),
    ("Dashboard", "React UI"),
]

for i, (title, desc) in enumerate(flow_steps):
    x = Inches(0.5 + i * 2.5)
    box = add_shape_box(slide3, x, Inches(4.3), Inches(2.2), Inches(1.2), RGBColor(0x1e, 0x29, 0x3b), PRIMARY)
    add_textbox(slide3, x + Inches(0.1), Inches(4.4), Inches(2.0), Inches(0.3),
                title, font_size=13, bold=True, color=PRIMARY)
    add_textbox(slide3, x + Inches(0.1), Inches(4.7), Inches(2.0), Inches(0.7),
                desc, font_size=11, color=GRAY)
    if i < 4:
        add_textbox(slide3, x + Inches(2.2), Inches(4.6), Inches(0.3), Inches(0.3),
                    "→", font_size=18, bold=True, color=PRIMARY)

# Methodology
add_textbox(slide3, Inches(0.5), Inches(5.7), Inches(12), Inches(0.4),
            "Implementation Methodology", font_size=16, bold=True, color=WHITE)
method = [
    "1. Data Integration → Real-time sync from TMS, SMMS, TDMS, COA",
    "2. AI Optimization → Genetic Algorithm for multi-objective scheduling",
    "3. Constraint Engine → Time windows, resources, department coordination",
    "4. Plan Generation → Weekly + Monthly block schedules",
    "5. Dashboard → Interactive calendar with AI recommendations"
]
add_bullet_text(slide3, Inches(0.5), Inches(6.1), Inches(12), Inches(1.2), method, font_size=12, color=GRAY)

# ============================================
# SLIDE 4: FEASIBILITY AND VIABILITY
# ============================================
slide4 = prs.slides.add_slide(prs.slide_layouts[6])
add_background(slide4, DARK)

add_textbox(slide4, Inches(0.5), Inches(0.3), Inches(12), Inches(0.6),
            "FEASIBILITY AND VIABILITY", font_size=28, bold=True, color=PRIMARY)

# Feasibility
box3 = add_shape_box(slide4, Inches(0.5), Inches(1.1), Inches(6), Inches(3.0), RGBColor(0x1e, 0x29, 0x3b))
add_textbox(slide4, Inches(0.8), Inches(1.2), Inches(5.5), Inches(0.4),
            "Feasibility Analysis", font_size=16, bold=True, color=GREEN)
feasibility = [
    "BDMS system already deployed on Central Railway",
    "TMS, SMMS, TDMS APIs available via CRIS",
    "COA corridor data accessible",
    "AI/ML algorithms proven in similar domains",
    "Scalable cloud infrastructure available"
]
add_bullet_text(slide4, Inches(0.8), Inches(1.7), Inches(5.5), Inches(2.2), feasibility, font_size=13, color=GRAY)

# Challenges
box4 = add_shape_box(slide4, Inches(6.8), Inches(1.1), Inches(6), Inches(3.0), RGBColor(0x1e, 0x29, 0x3b))
add_textbox(slide4, Inches(7.1), Inches(1.2), Inches(5.5), Inches(0.4),
            "Potential Challenges", font_size=16, bold=True, color=ORANGE)
challenges = [
    "Data quality and consistency across systems",
    "Legacy system integration complexity",
    "User adoption and training requirements",
    "Real-time synchronization latency",
    "Scalability across all railway zones"
]
add_bullet_text(slide4, Inches(7.1), Inches(1.7), Inches(5.5), Inches(2.2), challenges, font_size=13, color=GRAY)

# Strategies
box5 = add_shape_box(slide4, Inches(0.5), Inches(4.3), Inches(12.3), Inches(3.0), RGBColor(0x1e, 0x29, 0x3b))
add_textbox(slide4, Inches(0.8), Inches(4.4), Inches(11.8), Inches(0.4),
            "Strategies for Overcoming Challenges", font_size=16, bold=True, color=PRIMARY)
strategies = [
    "Phase 1: Pilot on Mumbai Division (already has BDMS)",
    "Phase 2: Data validation layer with automated quality checks",
    "Phase 3: REST API wrappers for legacy system integration",
    "Phase 4: User training program with field staff feedback",
    "Phase 5: Gradual rollout to other divisions with monitoring"
]
add_bullet_text(slide4, Inches(0.8), Inches(4.9), Inches(11.8), Inches(2.2), strategies, font_size=13, color=GRAY)

# ============================================
# SLIDE 5: IMPACT AND BENEFITS
# ============================================
slide5 = prs.slides.add_slide(prs.slide_layouts[6])
add_background(slide5, DARK)

add_textbox(slide5, Inches(0.5), Inches(0.3), Inches(12), Inches(0.6),
            "IMPACT AND BENEFITS", font_size=28, bold=True, color=PRIMARY)

# Impact metrics
metrics = [
    ("30%", "Increase in\nAsset Availability", PRIMARY),
    ("40%", "Faster Block\nApproval Process", GREEN),
    ("25%", "Reduction in\nMaintenance Delays", ORANGE),
    ("Zero", "Train Service\nDisruption", RED),
]

for i, (value, label, color) in enumerate(metrics):
    x = Inches(0.5 + i * 3.2)
    box = add_shape_box(slide5, x, Inches(1.1), Inches(2.9), Inches(1.5), RGBColor(0x1e, 0x29, 0x3b), color)
    add_textbox(slide5, x + Inches(0.2), Inches(1.2), Inches(2.5), Inches(0.6),
                value, font_size=32, bold=True, color=color, alignment=PP_ALIGN.CENTER)
    add_textbox(slide5, x + Inches(0.2), Inches(1.8), Inches(2.5), Inches(0.7),
                label, font_size=13, color=GRAY, alignment=PP_ALIGN.CENTER)

# Benefits
add_textbox(slide5, Inches(0.5), Inches(2.8), Inches(12), Inches(0.4),
            "Benefits", font_size=18, bold=True, color=WHITE)

benefits = [
    ("Social", "Enhanced safety for passengers and railway staff through better maintenance planning", GREEN),
    ("Economic", "Reduced downtime costs, optimized resource utilization, fewer train cancellations", PRIMARY),
    ("Operational", "Coordinated multi-department scheduling, real-time visibility, audit trail", ORANGE),
]

for i, (title, desc, color) in enumerate(benefits):
    y = Inches(3.3 + i * 1.0)
    box = add_shape_box(slide5, Inches(0.5), y, Inches(12.3), Inches(0.9), RGBColor(0x1e, 0x29, 0x3b), color)
    add_textbox(slide5, Inches(0.8), y + Inches(0.1), Inches(2), Inches(0.3),
                title, font_size=14, bold=True, color=color)
    add_textbox(slide5, Inches(0.8), y + Inches(0.4), Inches(11.5), Inches(0.4),
                desc, font_size=13, color=GRAY)

# Target Audience
add_textbox(slide5, Inches(0.5), Inches(6.3), Inches(12), Inches(0.4),
            "Target Audience: Indian Railways (Engineering, Traction, S&T Departments), Railway Board, CRIS",
            font_size=13, bold=True, color=PRIMARY)

# ============================================
# SLIDE 6: RESEARCH AND REFERENCES
# ============================================
slide6 = prs.slides.add_slide(prs.slide_layouts[6])
add_background(slide6, DARK)

add_textbox(slide6, Inches(0.5), Inches(0.3), Inches(12), Inches(0.6),
            "RESEARCH AND REFERENCES", font_size=28, bold=True, color=PRIMARY)

# References
add_textbox(slide6, Inches(0.5), Inches(1.0), Inches(12), Inches(0.4),
            "Research Papers", font_size=16, bold=True, color=WHITE)

refs = [
    "1. Huang et al. (2026) - Robust RL with GNN for stochastic urban rail maintenance scheduling",
    "2. Yang et al. (2026) - Robust rescheduling of train timetables and maintenance windows",
    "3. Arcieri et al. (2025) - Graph-based multi-agent RL for railway infrastructure decision support",
    "4. Mehranfar et al. (2025) - Optimizing multicomponent intervention programs",
    "5. Zhang et al. (2020) - Joint optimization of train scheduling and maintenance planning",
]
add_bullet_text(slide6, Inches(0.5), Inches(1.5), Inches(12), Inches(2.5), refs, font_size=12, color=GRAY)

# Existing Systems
add_textbox(slide6, Inches(0.5), Inches(3.8), Inches(12), Inches(0.4),
            "Existing Indian Railways Systems Referenced", font_size=16, bold=True, color=WHITE)

systems = [
    "BDMS (Block & Disconnection Management System) - CRIS",
    "TMS (Track Management System)",
    "SMMS (Signalling Maintenance & Management System)",
    "TDMS (Traction Distribution Management System)",
    "COA (Control Office Application)",
    "Rolling Block Programme - Railway Board Guidelines"
]
add_bullet_text(slide6, Inches(0.5), Inches(4.3), Inches(12), Inches(2.5), systems, font_size=13, color=GRAY)

# Thank you
add_textbox(slide6, Inches(0.5), Inches(6.5), Inches(12), Inches(0.5),
            "Thank You | Niravaan | Team [Your Team Name]",
            font_size=16, bold=True, color=PRIMARY, alignment=PP_ALIGN.CENTER)

# Save
output_path = r'C:\Users\Abhinav\sih-railways\presentation\SIH26027_Niravaan_Presentation.pptx'
prs.save(output_path)
print(f"Presentation saved to: {output_path}")
